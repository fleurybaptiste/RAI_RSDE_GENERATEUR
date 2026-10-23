# -*- coding: utf-8 -*-
"""
Convertit l'onglet 'Resultats' d'un classeur de saisie SANDRE FCT_ASSAIN V4
(RSDE) en fichier XML conforme au scenario d'echange - version avec fenetre.
"""
import os
import sys
import traceback
import threading
import xml.etree.ElementTree as ET
from xml.dom import minidom

import tkinter as tk
from tkinter import filedialog, messagebox
from tkinter import ttk

try:
    import openpyxl
except ImportError:
    import tkinter.messagebox as mb
    root = tk.Tk()
    root.withdraw()
    mb.showerror("Erreur", "Le module openpyxl est introuvable. Reinstallez l'outil.")
    sys.exit(1)

NS = "http://xml.sandre.eaufrance.fr/scenario/fct_assain/4"
XSI = "http://www.w3.org/2001/XMLSchema-instance"
SCHEMA_LOCATION = (
    "http://xml.sandre.eaufrance.fr/scenario/fct_assain/4 "
    "http://xml.sandre.eaufrance.fr/scenario/fct_assain/4/sandre_sc_fct_assain.xsd"
)

COLUMNS = [
    "SIRET_Emetteur", "Nom_Emetteur", "SIRET_Destinataire", "Nom_Destinataire",
    "DateDebutReference", "DateFinReference",
    "CdOuvrageDepollution", "TypeOuvrageDepollution", "NomOuvrageDepollution",
    "CdCommune", "LbCommune", "NatureSystTraitementEauxUsees",
    "SIRET_MaitreOuvrage", "Nom_MaitreOuvrage", "SIRET_Exploitant", "Nom_Exploitant",
    "NumeroPointMesure", "LbPointMesure", "LocGlobalePointMesure", "ComPointMesure",
    "DatePrlvt", "HeurePrlvt", "DureePrlvt", "SIRET_Preleveur", "Nom_Preleveur",
    "ConformitePrlvt", "AccrePrlvt", "CdSupport",
    "DateReceptionEchantillon", "HeureReceptionEchantillon", "DateAnalyse", "HeureAnalyse",
    "CdParametre", "RsAnalyse", "CdUniteMesure", "CdRemAnalyse", "InSituAnalyse",
    "StatutRsAnalyse", "QualRsAnalyse", "CdFractionAnalysee", "CdMethode",
    "SIRET_Laboratoire", "Nom_Laboratoire", "SIRET_Producteur", "Nom_Producteur",
    "FinaliteAnalyse", "LDAna", "LQAna", "LSAna", "AccreAna", "AgreAna", "IncertAna", "ComAna",
]

OBLIGATOIRES = {
    "SIRET_Emetteur", "Nom_Emetteur", "DateDebutReference", "DateFinReference",
    "CdOuvrageDepollution", "TypeOuvrageDepollution", "NomOuvrageDepollution",
    "NatureSystTraitementEauxUsees",
    "NumeroPointMesure", "LbPointMesure", "LocGlobalePointMesure",
    "DatePrlvt", "HeurePrlvt", "DureePrlvt", "SIRET_Preleveur", "ConformitePrlvt",
    "AccrePrlvt", "CdSupport",
    "DateReceptionEchantillon", "HeureReceptionEchantillon", "DateAnalyse",
    "CdParametre", "RsAnalyse", "CdUniteMesure", "CdRemAnalyse", "InSituAnalyse",
    "StatutRsAnalyse", "QualRsAnalyse", "CdFractionAnalysee", "CdMethode",
    "SIRET_Laboratoire", "FinaliteAnalyse", "LQAna", "AccreAna", "AgreAna", "IncertAna",
}


def cellstr(v):
    if v is None:
        return ""
    return str(v).strip()


def read_rows(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if "Resultats" not in wb.sheetnames:
        raise ValueError("L'onglet 'Resultats' est introuvable dans ce classeur.")
    ws = wb["Resultats"]

    data_start = 3
    for r in range(1, 6):
        a1 = str(ws.cell(row=r, column=1).value or "")
        if "SIRET_Emetteur" in a1:
            data_start = r + 1
            break

    rows = []
    r = data_start
    while True:
        any_val = any(ws.cell(row=r, column=c).value not in (None, "") for c in range(1, len(COLUMNS) + 1))
        if not any_val:
            break
        row = {}
        for ci, name in enumerate(COLUMNS, start=1):
            row[name] = cellstr(ws.cell(row=r, column=ci).value)
        row["_excel_row"] = r
        rows.append(row)
        r += 1
    return rows


def check_missing(rows):
    problems = []
    for row in rows:
        r = row["_excel_row"]
        for col in OBLIGATOIRES:
            if row.get(col, "") == "":
                if col == "RsAnalyse" and row.get("CdRemAnalyse") == "0":
                    continue
                problems.append((r, col))
    return problems


def sub(parent, tag, text=None, attrib=None):
    el = ET.SubElement(parent, tag, attrib or {})
    if text is not None and text != "":
        el.text = text
    return el


def build_intervenant(parent, tag, siret, nom):
    if not siret and not nom:
        return None
    el = sub(parent, tag)
    sub(el, "CdIntervenant", siret, {"schemeAgencyID": "SIRET"})
    if nom:
        sub(el, "NomIntervenant", nom)
    return el


def build_xml(rows):
    ET.register_namespace("", NS)
    root = ET.Element("FctAssain", {
        "xmlns": NS,
        "xmlns:xsi": XSI,
        "xsi:schemaLocation": SCHEMA_LOCATION,
    })

    first = rows[0]
    scenario = sub(root, "Scenario")
    sub(scenario, "CodeScenario", "FCT_ASSAIN")
    sub(scenario, "VersionScenario", "4")
    sub(scenario, "NomScenario",
        "Autosurveillance des systèmes de collecte et de traitement des eaux usées d'origine urbaine")
    sub(scenario, "DateDebutReference", first["DateDebutReference"])
    sub(scenario, "DateFinReference", first["DateFinReference"])
    build_intervenant(scenario, "Emetteur", first["SIRET_Emetteur"], first["Nom_Emetteur"])
    if first["SIRET_Destinataire"] or first["Nom_Destinataire"]:
        build_intervenant(scenario, "Destinataire", first["SIRET_Destinataire"], first["Nom_Destinataire"])

    ouvrages = {}
    for row in rows:
        cd_ouv = row["CdOuvrageDepollution"]
        ouvrages.setdefault(cd_ouv, {"info": row, "points": {}})
        pts = ouvrages[cd_ouv]["points"]
        key_pt = (row["NumeroPointMesure"], row["LbPointMesure"], row["LocGlobalePointMesure"])
        pts.setdefault(key_pt, {"info": row, "prlvts": {}})
        prlvts = pts[key_pt]["prlvts"]
        key_pr = (row["DatePrlvt"], row["CdSupport"])
        prlvts.setdefault(key_pr, {"info": row, "analyses": []})
        prlvts[key_pr]["analyses"].append(row)

    for cd_ouv, ouv_data in ouvrages.items():
        info = ouv_data["info"]
        ouv_el = sub(root, "OuvrageDepollution")
        sub(ouv_el, "CdOuvrageDepollution", cd_ouv)
        sub(ouv_el, "TypeOuvrageDepollution", info["TypeOuvrageDepollution"])
        sub(ouv_el, "NomOuvrageDepollution", info["NomOuvrageDepollution"])
        if info["CdCommune"] or info["LbCommune"]:
            commune = sub(ouv_el, "Commune")
            sub(commune, "CdCommune", info["CdCommune"])
            if info["LbCommune"]:
                sub(commune, "LbCommune", info["LbCommune"])
        sub(ouv_el, "NatureSystTraitementEauxUsees", info["NatureSystTraitementEauxUsees"])
        if info["SIRET_MaitreOuvrage"] or info["Nom_MaitreOuvrage"]:
            build_intervenant(ouv_el, "MaitreDOuvrage", info["SIRET_MaitreOuvrage"], info["Nom_MaitreOuvrage"])
        if info["SIRET_Exploitant"] or info["Nom_Exploitant"]:
            build_intervenant(ouv_el, "Exploitant", info["SIRET_Exploitant"], info["Nom_Exploitant"])

        for key_pt, pt_data in ouv_data["points"].items():
            num_pm, lb_pm, loc_pm = key_pt
            pt_info = pt_data["info"]
            pt_el = sub(ouv_el, "PointMesure")
            sub(pt_el, "NumeroPointMesure", num_pm)
            sub(pt_el, "LbPointMesure", lb_pm)
            sub(pt_el, "LocGlobalePointMesure", loc_pm)
            if pt_info["ComPointMesure"]:
                sub(pt_el, "ComPointMesure", pt_info["ComPointMesure"])

            for key_pr, pr_data in pt_data["prlvts"].items():
                date_pr, cd_support = key_pr
                pr_info = pr_data["info"]
                prlvt_el = sub(pt_el, "Prlvt")
                sub(prlvt_el, "DatePrlvt", date_pr)
                if pr_info["HeurePrlvt"]:
                    sub(prlvt_el, "HeurePrlvt", pr_info["HeurePrlvt"])
                if pr_info["DureePrlvt"]:
                    sub(prlvt_el, "DureePrlvt", pr_info["DureePrlvt"])
                if pr_info["SIRET_Preleveur"] or pr_info["Nom_Preleveur"]:
                    build_intervenant(prlvt_el, "Preleveur", pr_info["SIRET_Preleveur"], pr_info["Nom_Preleveur"])
                if pr_info["ConformitePrlvt"]:
                    sub(prlvt_el, "ConformitePrlvt", pr_info["ConformitePrlvt"])
                if pr_info["AccrePrlvt"]:
                    sub(prlvt_el, "AccrePrlvt", pr_info["AccrePrlvt"])
                support_el = sub(prlvt_el, "Support")
                sub(support_el, "CdSupport", cd_support)

                for a in pr_data["analyses"]:
                    ana_el = sub(prlvt_el, "Analyse")
                    if a["DateReceptionEchantillon"]:
                        sub(ana_el, "DateReceptionEchantillon", a["DateReceptionEchantillon"])
                    if a["HeureReceptionEchantillon"]:
                        sub(ana_el, "HeureReceptionEchantillon", a["HeureReceptionEchantillon"])
                    if a["DateAnalyse"]:
                        sub(ana_el, "DateAnalyse", a["DateAnalyse"])
                    if a["HeureAnalyse"]:
                        sub(ana_el, "HeureAnalyse", a["HeureAnalyse"])
                    sub(ana_el, "RsAnalyse", a["RsAnalyse"])
                    sub(ana_el, "CdRemAnalyse", a["CdRemAnalyse"])
                    sub(ana_el, "InSituAnalyse", a["InSituAnalyse"])
                    sub(ana_el, "StatutRsAnalyse", a["StatutRsAnalyse"])
                    sub(ana_el, "QualRsAnalyse", a["QualRsAnalyse"])
                    frac_el = sub(ana_el, "FractionAnalysee")
                    sub(frac_el, "CdFractionAnalysee", a["CdFractionAnalysee"])
                    if a["CdMethode"]:
                        meth_el = sub(ana_el, "MethodeAna")
                        sub(meth_el, "CdMethode", a["CdMethode"])
                    param_el = sub(ana_el, "Parametre")
                    sub(param_el, "CdParametre", a["CdParametre"])
                    unite_el = sub(ana_el, "UniteMesure")
                    sub(unite_el, "CdUniteMesure", a["CdUniteMesure"])
                    if a["SIRET_Laboratoire"] or a["Nom_Laboratoire"]:
                        build_intervenant(ana_el, "Laboratoire", a["SIRET_Laboratoire"], a["Nom_Laboratoire"])
                    if a["SIRET_Producteur"] or a["Nom_Producteur"]:
                        build_intervenant(ana_el, "Producteur", a["SIRET_Producteur"], a["Nom_Producteur"])
                    sub(ana_el, "FinaliteAnalyse", a["FinaliteAnalyse"])
                    if a["LDAna"]:
                        sub(ana_el, "LDAna", a["LDAna"])
                    if a["LQAna"]:
                        sub(ana_el, "LQAna", a["LQAna"])
                    if a["LSAna"]:
                        sub(ana_el, "LSAna", a["LSAna"])
                    if a["AccreAna"]:
                        sub(ana_el, "AccreAna", a["AccreAna"])
                    if a["AgreAna"]:
                        sub(ana_el, "AgreAna", a["AgreAna"])
                    if a["IncertAna"]:
                        sub(ana_el, "IncertAna", a["IncertAna"])
                    if a["ComAna"]:
                        sub(ana_el, "ComAna", a["ComAna"])

    return root


def prettify(root):
    rough = ET.tostring(root, encoding="utf-8")
    reparsed = minidom.parseString(rough)
    pretty = reparsed.toprettyxml(indent="  ", encoding="UTF-8")
    lines = [l for l in pretty.decode("utf-8").split("\n") if l.strip()]
    return "\n".join(lines)


def convert(xlsx_path, log):
    log(f"Lecture de : {xlsx_path}")
    rows = read_rows(xlsx_path)
    if not rows:
        log("Aucune ligne de donnees trouvee dans l'onglet 'Resultats'.")
        return
    log(f"{len(rows)} ligne(s) de resultat trouvee(s).")

    problems = check_missing(rows)
    base, _ = os.path.splitext(xlsx_path)
    report_path = base + "_rapport.txt"
    with open(report_path, "w", encoding="utf-8") as f:
        if problems:
            f.write(f"{len(problems)} champ(s) obligatoire(s) manquant(s) :\n")
            for r, col in problems:
                f.write(f"  ligne Excel {r} : colonne '{col}' vide\n")
        else:
            f.write("Aucun champ obligatoire manquant detecte.\n")
        f.write("\nRappel : ce controle ne remplace pas la validation officielle Sandre\n")
        f.write("(www.sandre.eaufrance.fr) a effectuer avant tout envoi.\n")

    if problems:
        log(f"\n/!\\ {len(problems)} champ(s) obligatoire(s) manquant(s) - voir {os.path.basename(report_path)}")
        for r, col in problems[:30]:
            log(f"  ligne Excel {r} : colonne '{col}' vide")
        if len(problems) > 30:
            log(f"  ... et {len(problems) - 30} autre(s), voir le rapport complet.")
    else:
        log("Aucun champ obligatoire manquant detecte.")

    log("\nGeneration du fichier XML...")
    root = build_xml(rows)
    xml_text = prettify(root)

    xml_path = base + ".xml"
    with open(xml_path, "w", encoding="utf-8") as f:
        f.write(xml_text)

    log(f"\nFichier XML genere : {xml_path}")
    log(f"Rapport de controle : {report_path}")
    log("\nRappel : validez ce fichier sur www.sandre.eaufrance.fr avant envoi,")
    log("et compressez-le au format GZip selon la regle de nommage du scenario.")
    return xml_path, report_path, len(problems)


class App:
    def __init__(self, root):
        self.root = root
        root.title("SANDRE - Conversion classeur vers XML")
        root.geometry("720x520")
        root.minsize(600, 420)

        pad = {"padx": 10, "pady": 8}

        top = ttk.Frame(root)
        top.pack(fill="x", **pad)

        ttk.Label(top, text="Classeur de saisie (.xlsx) :").pack(side="left")
        self.path_var = tk.StringVar()
        entry = ttk.Entry(top, textvariable=self.path_var)
        entry.pack(side="left", fill="x", expand=True, padx=(8, 8))
        ttk.Button(top, text="Parcourir...", command=self.browse).pack(side="left")

        action = ttk.Frame(root)
        action.pack(fill="x", padx=10)
        self.convert_btn = ttk.Button(action, text="Générer le XML", command=self.on_convert)
        self.convert_btn.pack(side="left")
        self.open_folder_btn = ttk.Button(action, text="Ouvrir le dossier du résultat",
                                           command=self.open_folder, state="disabled")
        self.open_folder_btn.pack(side="left", padx=(8, 0))

        self.status_var = tk.StringVar(value="En attente d'un fichier...")
        ttk.Label(root, textvariable=self.status_var).pack(fill="x", padx=10, pady=(4, 0))

        log_frame = ttk.Frame(root)
        log_frame.pack(fill="both", expand=True, padx=10, pady=10)
        self.log_text = tk.Text(log_frame, wrap="word", state="disabled", bg="#f5f5f5")
        scroll = ttk.Scrollbar(log_frame, command=self.log_text.yview)
        self.log_text.configure(yscrollcommand=scroll.set)
        self.log_text.pack(side="left", fill="both", expand=True)
        scroll.pack(side="right", fill="y")

        self.last_output_dir = None

    def log(self, msg):
        self.log_text.configure(state="normal")
        self.log_text.insert("end", str(msg) + "\n")
        self.log_text.see("end")
        self.log_text.configure(state="disabled")
        self.root.update_idletasks()

    def browse(self):
        path = filedialog.askopenfilename(
            title="Choisir le classeur de saisie SANDRE (.xlsx)",
            filetypes=[("Classeur Excel", "*.xlsx")],
            parent=self.root,
        )
        if path:
            self.path_var.set(path)

    def open_folder(self):
        if self.last_output_dir:
            os.startfile(self.last_output_dir)

    def on_convert(self):
        path = self.path_var.get().strip()
        if not path:
            messagebox.showwarning("Fichier manquant", "Choisissez d'abord un classeur .xlsx.", parent=self.root)
            return
        if not os.path.isfile(path):
            messagebox.showerror("Fichier introuvable", f"Le fichier n'existe pas :\n{path}", parent=self.root)
            return

        self.convert_btn.configure(state="disabled")
        self.open_folder_btn.configure(state="disabled")
        self.status_var.set("Conversion en cours...")
        self.log_text.configure(state="normal")
        self.log_text.delete("1.0", "end")
        self.log_text.configure(state="disabled")

        thread = threading.Thread(target=self._run_conversion, args=(path,), daemon=True)
        thread.start()

    def _run_conversion(self, path):
        try:
            result = convert(path, self.log)
        except Exception as e:
            self.log(f"\nERREUR : {e}")
            self.log(traceback.format_exc())
            self.root.after(0, lambda: self._on_done(False, None))
            return
        self.root.after(0, lambda: self._on_done(True, result))

    def _on_done(self, success, result):
        self.convert_btn.configure(state="normal")
        if success and result:
            xml_path, report_path, nb_problems = result
            self.last_output_dir = os.path.dirname(xml_path)
            self.open_folder_btn.configure(state="normal")
            if nb_problems:
                self.status_var.set(f"Terminé avec {nb_problems} champ(s) obligatoire(s) manquant(s).")
            else:
                self.status_var.set("Terminé : aucun champ obligatoire manquant.")
            messagebox.showinfo(
                "Conversion terminée",
                f"Fichier XML généré :\n{xml_path}\n\n"
                + (f"{nb_problems} champ(s) obligatoire(s) manquant(s), voir le rapport."
                   if nb_problems else "Aucun champ obligatoire manquant."),
                parent=self.root,
            )
        else:
            self.status_var.set("Échec de la conversion, voir le journal ci-dessus.")
            messagebox.showerror("Erreur", "La conversion a échoué, voir le détail dans la fenêtre.", parent=self.root)


def main():
    root = tk.Tk()
    try:
        style = ttk.Style()
        if "vista" in style.theme_names():
            style.theme_use("vista")
    except Exception:
        pass
    app = App(root)
    if len(sys.argv) > 1 and os.path.isfile(sys.argv[1]):
        app.path_var.set(sys.argv[1])
    root.mainloop()


if __name__ == "__main__":
    main()
